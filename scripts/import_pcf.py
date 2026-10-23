"""把 APQC PCF v8.0 的 Excel 文件转成可入库的 JSON 数据。

用法：

    python scripts/import_pcf.py \
        --cross-industry  "K016808_..._Cross-Industry...xlsx" \
        --consumer-products "K017170_..._Consumer Products...xlsx" \
        --consumer-electronics "K017180_..._Consumer Electronics...xlsx" \
        --chinese "APQC_PCF_v80_跨行业流程分类框架_中文整理版.xlsx" \
        --output app/seed/data/pcf_v80.json

设计口径：

- **PCF ID 是跨版本稳定键**，`Hierarchy ID` 会因行业元素插入而漂移，不能当外键；
- 每个变体保留自己官方的 Hierarchy ID，父节点按同文件内的层级路径解析；
- 中文名只覆盖 L1–L3（449 个元素），按 PCF ID 与英文名双向校验后合并；
- L4/L5 没有中文来源，`name_zh` 记为空，展示时回落英文；
- 指标按 `(PCF ID, Metric ID)` 去重，跨变体合并成一张指标目录。

脚本只生成数据，不连数据库；产物由 `app/seed/pcf_data.py` 读取。
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import openpyxl

# 变体定义。code 与现有对标框架命名保持一致；APQC_PCF 是主框架。
FRAMEWORK_SPECS: List[Dict[str, Any]] = [
    {
        "code": "APQC_PCF",
        "name": "APQC 流程分类框架（跨行业）",
        "name_en": "APQC Process Classification Framework",
        "industry": "CROSS_INDUSTRY",
        "is_primary": True,
        "base_framework_code": None,
        "description": "APQC PCF v8.0 跨行业版，作为系统默认的流程分类主框架。",
        "sort_order": 2,
        "source_key": "cross_industry",
    },
    {
        "code": "APQC_PCF_CP",
        "name": "APQC 流程分类框架（消费品）",
        "name_en": "APQC Process Classification Framework - Consumer Products",
        "industry": "CONSUMER_PRODUCTS",
        "is_primary": False,
        "base_framework_code": "APQC_PCF",
        "description": "APQC PCF v8.0 消费品行业版，基于跨行业版增加行业专属流程。",
        "sort_order": 3,
        "source_key": "consumer_products",
    },
    {
        "code": "APQC_PCF_CE",
        "name": "APQC 流程分类框架（消费电子）",
        "name_en": "APQC Process Classification Framework - Consumer Electronics",
        "industry": "CONSUMER_ELECTRONICS",
        "is_primary": False,
        "base_framework_code": "APQC_PCF",
        "description": "APQC PCF v8.0 消费电子行业版，基于跨行业版增加行业专属流程。",
        "sort_order": 4,
        "source_key": "consumer_electronics",
    },
]


class ImportError_(ValueError):
    """源文件不符合预期结构时抛出，便于脚本给出明确错误。"""


# ------------------------------------------------------------------ 基础工具


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header_index(ws, required: Iterable[str]) -> Dict[str, int]:
    """按首行表头建立索引，忽略大小写、空格和换行。"""
    headers = {}
    for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))):
        key = _text(cell).lower().replace("\n", "").replace(" ", "")
        if key:
            headers[key] = index
    required_keys = {name: name.lower().replace("\n", "").replace(" ", "") for name in required}
    missing = [name for name, key in required_keys.items() if key not in headers]
    if missing:
        raise ImportError_(f"{ws.title} 缺少表头：{', '.join(missing)}")
    return {name: headers[key] for name, key in required_keys.items()}


def _level(hierarchy: str) -> int:
    """把 1.0 / 1.1 / 1.1.1 这类编号解析成层级。"""
    parts = hierarchy.split(".")
    if len(parts) < 2 or not all(part.isdigit() for part in parts):
        raise ImportError_(f"无法解析 Hierarchy ID：{hierarchy}")
    if len(parts) == 2 and parts[1] == "0":
        return 1
    return len(parts)


def _parent_hierarchy(hierarchy: str, level: int) -> Optional[str]:
    parts = hierarchy.split(".")
    if level == 1:
        return None
    if level == 2:
        return f"{parts[0]}.0"
    return ".".join(parts[:-1])


def _normalize_name(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def _iso_date(value: str) -> Optional[str]:
    match = re.search(r"generated on\s+(\d{1,2}/\d{1,2}/\d{4})", value, re.IGNORECASE)
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%m/%d/%Y").date().isoformat()
    except ValueError:
        return None


def _attribution(wb) -> Optional[str]:
    """提取版权/署名段落，随框架一起入库，满足 APQC 的署名要求。"""
    if "Copyright and Attribution" not in wb.sheetnames:
        return None
    parts = [
        _text(cell)
        for row in wb["Copyright and Attribution"].iter_rows(values_only=True)
        for cell in row
        if _text(cell)
    ]
    text = "\n".join(parts)
    marker = text.upper().find("COPYRIGHT AND ATTRIBUTION")
    if marker >= 0:
        text = text[marker:]
    return text[:4000] or None


# ------------------------------------------------------------------ 读取源文件


def load_chinese_names(path: Path) -> Dict[str, Dict[str, Any]]:
    """读取中文整理版的 L1–L3，返回 PCF ID -> 中英文名。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "02_L1-L3流程清单" not in wb.sheetnames:
        raise ImportError_(f"{path.name} 缺少工作表 02_L1-L3流程清单")
    ws = wb["02_L1-L3流程清单"]
    columns = _header_index(ws, ["层级", "Hierarchy ID", "PCF ID", "英文名", "中文名"])

    result: Dict[str, Dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        level = _text(row[columns["层级"]])
        external_id = _text(row[columns["PCF ID"]])
        name_en = _text(row[columns["英文名"]])
        name_zh = _text(row[columns["中文名"]])
        hierarchy = _text(row[columns["Hierarchy ID"]])
        if not external_id or not name_zh or not hierarchy:
            continue
        expected_level = int(level.lstrip("Ll")) if level else _level(hierarchy)
        if expected_level > 3:
            continue
        if external_id in result and _normalize_name(
            result[external_id]["name_en"]
        ) != _normalize_name(name_en):
            raise ImportError_(f"中文表里 PCF ID {external_id} 出现冲突的英文名")
        result[external_id] = {
            "name_zh": name_zh,
            "name_en": name_en,
            "level": expected_level,
            "hierarchy": hierarchy,
        }
    if not result:
        raise ImportError_(f"{path.name} 没有解析到任何中文名称")
    return result


def load_variant(
    path: Path, spec: Dict[str, Any], zh_map: Dict[str, Dict[str, Any]]
) -> Dict[str, Any]:
    """读取一个 PCF v8.0 变体，返回带框架元数据的分类树。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Combined" not in wb.sheetnames:
        raise ImportError_(f"{path.name} 缺少 Combined 工作表")
    ws = wb["Combined"]
    columns = _header_index(
        ws,
        [
            "PCF ID",
            "Hierarchy ID",
            "Name",
            "Difference Index",
            "Change Details",
            "Metrics Available?",
            "Element Description",
        ],
    )

    raw_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        external_id = _text(row[columns["PCF ID"]])
        hierarchy = _text(row[columns["Hierarchy ID"]])
        name_en = _text(row[columns["Name"]])
        if not external_id or not hierarchy or not name_en:
            continue
        level = _level(hierarchy)
        raw_rows.append(
            {
                "external_id": external_id,
                "code": hierarchy,
                "level": level,
                "name_en": name_en,
                "description": _text(row[columns["Element Description"]]) or None,
                "metrics_available": _text(row[columns["Metrics Available?"]]).upper() == "Y",
                "difference_index": _to_int(row[columns["Difference Index"]]),
                "change_details": _text(row[columns["Change Details"]]) or None,
            }
        )

    if not raw_rows:
        raise ImportError_(f"{path.name} 的 Combined 工作表没有数据")

    by_external = {row["external_id"]: row for row in raw_rows}
    if len(by_external) != len(raw_rows):
        raise ImportError_(f"{path.name} 存在重复的 PCF ID")
    by_hierarchy = {row["code"]: row for row in raw_rows}
    if len(by_hierarchy) != len(raw_rows):
        raise ImportError_(f"{path.name} 存在重复的 Hierarchy ID")

    # 父节点按同变体内的层级路径解析；父节点可以是行业版新增元素。
    for row in raw_rows:
        parent_hierarchy = _parent_hierarchy(row["code"], row["level"])
        if parent_hierarchy is None:
            row["parent_external_id"] = None
        else:
            parent = by_hierarchy.get(parent_hierarchy)
            if parent is None:
                raise ImportError_(f"{path.name} 中 {row['code']} 找不到父节点 {parent_hierarchy}")
            row["parent_external_id"] = parent["external_id"]

    # 中英文对照：英文名必须与中文表一致，避免编号漂移时错配。
    for row in raw_rows:
        zh = zh_map.get(row["external_id"])
        if zh is None:
            row["name_zh"] = None
            row["name"] = row["name_en"]
            continue
        if _normalize_name(zh["name_en"]) != _normalize_name(row["name_en"]):
            raise ImportError_(
                f"{path.name} 中 PCF ID {row['external_id']} 的英文名与中文表不一致："
                f"{row['name_en']!r} != {zh['name_en']!r}"
            )
        row["name_zh"] = zh["name_zh"]
        row["name"] = zh["name_zh"]

    # 同一父节点下按源文件顺序编号，保证树形展示稳定。
    counters: Dict[Optional[str], int] = {}
    for row in raw_rows:
        parent = row["parent_external_id"]
        row["sort_order"] = counters.get(parent, 0)
        counters[parent] = row["sort_order"] + 1

    counts: Dict[str, int] = {}
    for row in raw_rows:
        key = str(row["level"])
        counts[key] = counts.get(key, 0) + 1

    copyright_text = (
        "\n".join(
            _text(cell)
            for row in wb["Copyright and Attribution"].iter_rows(values_only=True)
            for cell in row
            if _text(cell)
        )
        if "Copyright and Attribution" in wb.sheetnames
        else ""
    )

    return {
        "code": spec["code"],
        "name": spec["name"],
        "name_en": spec["name_en"],
        "version": "8.0",
        "industry": spec["industry"],
        "is_primary": spec["is_primary"],
        "base_framework_code": spec["base_framework_code"],
        "description": spec["description"],
        "sort_order": spec["sort_order"],
        "source_file": path.name,
        "generated_on": _iso_date(copyright_text),
        "attribution": _attribution(wb),
        "counts": dict(sorted(counts.items())),
        "categories": raw_rows,
    }


def load_metrics(
    paths: Dict[str, Path], known_external_ids: set[str]
) -> Tuple[List[Dict[str, Any]], int]:
    """合并多个变体的指标，按 (PCF ID, Metric ID) 去重。

    跨变体同一指标的公式措辞可能不同；保留先读取的版本（主框架优先），
    并统计存在差异的指标键数量，便于追溯，不静默覆盖。
    """
    merged: Dict[Tuple[str, str], Dict[str, Any]] = {}
    conflict_keys: set[Tuple[str, str]] = set()
    for source_key, path in paths.items():
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Metrics" not in wb.sheetnames:
            continue
        ws = wb["Metrics"]
        columns = _header_index(
            ws,
            [
                "Process Element ID",
                "Metric Category",
                "Metric ID",
                "Metric Name",
                "Formula",
                "Units",
            ],
        )
        for row in ws.iter_rows(min_row=2, values_only=True):
            external_id = _text(row[columns["Process Element ID"]])
            metric_id = _text(row[columns["Metric ID"]])
            if not external_id or not metric_id:
                continue
            if external_id not in known_external_ids:
                continue
            key = (external_id, metric_id)
            payload = {
                "external_id": external_id,
                "metric_id": metric_id,
                "metric_category": _text(row[columns["Metric Category"]]) or None,
                "name": _text(row[columns["Metric Name"]]),
                "formula": _text(row[columns["Formula"]]) or None,
                "unit": _text(row[columns["Units"]]) or None,
            }
            previous = merged.get(key)
            if (
                previous is not None
                and {k: v for k, v in previous.items() if k != "source"} != payload
            ):
                conflict_keys.add(key)
                continue
            merged[key] = {**payload, "source": source_key}
    return [merged[key] for key in sorted(merged)], len(conflict_keys)


def _to_int(value: Any) -> int:
    text = _text(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


# ------------------------------------------------------------------ 主流程


def build_dataset(
    paths: Dict[str, Path],
    output: Path,
) -> Dict[str, Any]:
    zh_map = load_chinese_names(paths["chinese"])
    frameworks = [load_variant(paths[spec["source_key"]], spec, zh_map) for spec in FRAMEWORK_SPECS]

    known_external_ids = {
        row["external_id"] for framework in frameworks for row in framework["categories"]
    }
    metrics, conflicts = load_metrics(
        {
            key: paths[key]
            for key in ("cross_industry", "consumer_products", "consumer_electronics")
        },
        known_external_ids,
    )

    dataset = {
        "schema_version": 1,
        "pcf_version": "8.0",
        "generated_by": "scripts/import_pcf.py",
        "sources": {key: path.name for key, path in paths.items()},
        "chinese_coverage": {
            "elements": len(zh_map),
            "levels": sorted({row["level"] for row in zh_map.values()}),
        },
        "metrics": metrics,
        "metric_conflict_count": conflicts,
        "frameworks": frameworks,
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(dataset, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    return dataset


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="导入 APQC PCF v8.0 分类框架")
    parser.add_argument("--cross-industry", required=True, type=Path)
    parser.add_argument("--consumer-products", required=True, type=Path)
    parser.add_argument("--consumer-electronics", required=True, type=Path)
    parser.add_argument("--chinese", required=True, type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("app/seed/data/pcf_v80.json"),
        help="生成的 JSON 路径",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    paths = {
        "cross_industry": args.cross_industry,
        "consumer_products": args.consumer_products,
        "consumer_electronics": args.consumer_electronics,
        "chinese": args.chinese,
    }
    missing = [str(path) for path in paths.values() if not path.exists()]
    if missing:
        raise SystemExit("文件不存在：\n" + "\n".join(missing))

    dataset = build_dataset(paths, args.output)
    print(f"已写入 {args.output}")
    for framework in dataset["frameworks"]:
        counts = " / ".join(f"L{level}={count}" for level, count in framework["counts"].items())
        print(f"  {framework['code']}: {counts} 共 {len(framework['categories'])} 个元素")
    print(f"  中文覆盖: {dataset['chinese_coverage']['elements']} 个 L1–L3 元素")
    print(f"  指标: {len(dataset['metrics'])} 条（冲突 {dataset['metric_conflict_count']} 条）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
