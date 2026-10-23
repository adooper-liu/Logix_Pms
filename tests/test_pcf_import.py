"""APQC PCF v8.0 导入脚本与生成数据的回归测试。

覆盖三件事：
1. Excel 解析：层级、父子关系、按 PCF ID 合并中文、按变体保留 Hierarchy ID；
2. 指标合并：跨变体去重，公式差异只记录不覆盖；
3. 已提交的 `app/seed/data/pcf_v80.json` 完整性（数量、双语、父节点、指标）。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.seed import pcf_data
from scripts import import_pcf

COMBINED_HEADERS = [
    "PCF ID",
    "Hierarchy ID",
    "Name",
    "Difference Index",
    "Change Details",
    "Metrics Available?",
    "Element Description",
]
METRIC_HEADERS = [
    "Process Element ID",
    "Hierarchy ID",
    "Process Element",
    "Metric Category",
    "Metric ID",
    "Metric Name",
    "Formula",
    "Units",
]
CHINESE_HEADERS = ["层级", "Hierarchy ID", "PCF ID", "英文名", "中文名", "可对标", "指标数"]


def _write_variant(
    path: Path,
    rows: list[list[object]],
    metric_rows: list[list[object]] | None = None,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Combined"
    ws.append(COMBINED_HEADERS)
    for row in rows:
        ws.append(row)

    metrics = wb.create_sheet("Metrics")
    metrics.append(METRIC_HEADERS)
    for row in metric_rows or []:
        metrics.append(row)

    copyright_sheet = wb.create_sheet("Copyright and Attribution")
    copyright_sheet.append(["©2026 APQC. ALL RIGHTS RESERVED."])
    copyright_sheet.append(["This file was generated on 8/27/2026."])
    wb.save(path)


def _write_chinese(path: Path, rows: list[list[object]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "02_L1-L3流程清单"
    ws.append(CHINESE_HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)


@pytest.mark.parametrize(
    ("hierarchy", "level", "parent"),
    [
        ("1.0", 1, None),
        ("13.0", 1, None),
        ("1.1", 2, "1.0"),
        ("1.1.1", 3, "1.1"),
        ("1.1.1.1", 4, "1.1.1"),
        ("1.1.1.1.1", 5, "1.1.1.1"),
    ],
)
def test_level_and_parent(hierarchy, level, parent):
    assert import_pcf._level(hierarchy) == level
    assert import_pcf._parent_hierarchy(hierarchy, level) == parent


def test_build_dataset_merges_bilingual_names_and_dedupes_metrics(tmp_path):
    ci_rows = [
        ["100", "1.0", "Develop Vision and Strategy", "0", "", "Y", "Direction"],
        ["101", "1.1", "Define the business concept", "0", "", "N", "Concept"],
        ["102", "1.1.1", "Assess the external environment", "0", "", "Y", "Assess"],
        ["1021", "1.1.1.1", "Identify competitors", "0", "", "N", "Identify"],
    ]
    # 行业变体在 1.1 下插入新元素，导致后续编号漂移；PCF ID 仍稳定。
    cp_rows = ci_rows + [
        ["103", "1.1.2", "Consumer specific process", "0", "", "Y", "Consumer"],
    ]
    ce_rows = ci_rows + [
        ["104", "1.1.3", "Electronics specific process", "0", "", "N", "Electronics"],
    ]
    chinese_rows = [
        ["L1", "1.0", "100", "Develop Vision and Strategy", "制定愿景与战略", "Y", "2"],
        ["L2", "1.1", "101", "Define the business concept", "定义业务概念", "", "0"],
        ["L3", "1.1.1", "102", "Assess the external environment", "评估外部环境", "Y", "1"],
    ]

    paths = {
        "cross_industry": tmp_path / "ci.xlsx",
        "consumer_products": tmp_path / "cp.xlsx",
        "consumer_electronics": tmp_path / "ce.xlsx",
        "chinese": tmp_path / "zh.xlsx",
    }
    _write_variant(
        paths["cross_industry"],
        ci_rows,
        metric_rows=[
            [
                "100",
                "1.0",
                "Develop Vision and Strategy",
                "Process Efficiency",
                "M1",
                "Metric 1",
                "Formula A",
                "percent",
            ]
        ],
    )
    _write_variant(
        paths["consumer_products"],
        cp_rows,
        metric_rows=[
            [
                "103",
                "1.1.2",
                "Consumer specific process",
                "Cycle Time",
                "M2",
                "Metric 2",
                "Formula B",
                "days",
            ]
        ],
    )
    _write_variant(
        paths["consumer_electronics"],
        ce_rows,
        metric_rows=[
            # 同一指标在变体间公式措辞不同：保留主框架版本并记一次冲突。
            [
                "100",
                "1.0",
                "Develop Vision and Strategy",
                "Process Efficiency",
                "M1",
                "Metric 1",
                "Formula A2",
                "percent",
            ]
        ],
    )
    _write_chinese(paths["chinese"], chinese_rows)

    dataset = import_pcf.build_dataset(paths, tmp_path / "out.json")

    frameworks = {item["code"]: item for item in dataset["frameworks"]}
    assert set(frameworks) == {"APQC_PCF", "APQC_PCF_CP", "APQC_PCF_CE"}
    assert frameworks["APQC_PCF"]["is_primary"] is True
    assert frameworks["APQC_PCF_CP"]["base_framework_code"] == "APQC_PCF"
    assert frameworks["APQC_PCF"]["generated_on"] == "2026-08-27"

    ci = {row["external_id"]: row for row in frameworks["APQC_PCF"]["categories"]}
    assert ci["100"]["name_zh"] == "制定愿景与战略"
    assert ci["100"]["name"] == "制定愿景与战略"
    assert ci["100"]["name_en"] == "Develop Vision and Strategy"
    # L4 没有中文来源：name_zh 为空，name 回落英文。
    assert ci["1021"]["name_zh"] is None
    assert ci["1021"]["name"] == ci["1021"]["name_en"]
    assert ci["1021"]["parent_external_id"] == "102"

    cp = {row["external_id"]: row for row in frameworks["APQC_PCF_CP"]["categories"]}
    assert cp["103"]["parent_external_id"] == "101"

    # 指标按 (PCF ID, Metric ID) 去重：M1 保留主框架公式，M2 来自行业变体。
    metrics = {(item["external_id"], item["metric_id"]): item for item in dataset["metrics"]}
    assert set(metrics) == {("100", "M1"), ("103", "M2")}
    assert metrics[("100", "M1")]["formula"] == "Formula A"
    assert metrics[("100", "M1")]["source"] == "cross_industry"
    assert dataset["metric_conflict_count"] == 1

    written = tmp_path / "out.json"
    assert written.exists()
    assert "制定愿景与战略" in written.read_text(encoding="utf-8")


# ------------------------------------------------------------------ 已提交数据


@pytest.fixture(scope="module")
def dataset():
    return pcf_data.load_pcf_v80()


def test_committed_dataset_has_three_v8_frameworks(dataset):
    assert dataset["pcf_version"] == "8.0"
    assert [item["code"] for item in dataset["frameworks"]] == [
        "APQC_PCF",
        "APQC_PCF_CP",
        "APQC_PCF_CE",
    ]
    assert dataset["chinese_coverage"] == {"elements": 449, "levels": [1, 2, 3]}


def test_committed_framework_counts(dataset):
    expected = {
        "APQC_PCF": {1: 13, 2: 74, 3: 362, 4: 1353, 5: 215},
        "APQC_PCF_CP": {1: 13, 2: 74, 3: 365, 4: 1371, 5: 216},
        "APQC_PCF_CE": {1: 13, 2: 74, 3: 365, 4: 1379, 5: 216},
    }
    for framework in dataset["frameworks"]:
        counts = {}
        for row in framework["categories"]:
            counts[row["level"]] = counts.get(row["level"], 0) + 1
        assert counts == expected[framework["code"]], framework["code"]


def test_committed_tree_parents_resolve(dataset):
    for framework in dataset["frameworks"]:
        by_external = {row["external_id"]: row for row in framework["categories"]}
        for row in framework["categories"]:
            parent_id = row["parent_external_id"]
            if row["level"] == 1:
                assert parent_id is None
                continue
            parent = by_external.get(parent_id)
            assert parent is not None, (framework["code"], row["code"], parent_id)
            assert parent["level"] == row["level"] - 1


def test_committed_bilingual_pairs(dataset):
    """中英文都保留：L1–L3 有 name_zh，L4/L5 明确为空而不是覆盖英文。"""
    for framework in dataset["frameworks"]:
        chinese = [row for row in framework["categories"] if row["name_zh"]]
        assert len(chinese) == 449
        for row in framework["categories"]:
            assert row["name_en"]
            if row["name_zh"]:
                assert row["name"] == row["name_zh"]
            else:
                assert row["name"] == row["name_en"]

    sample = next(
        row for row in pcf_data.category_rows("APQC_PCF") if row["external_id"] == "10002"
    )
    assert sample["name_zh"] == "制定愿景与战略"
    assert sample["name_en"] == "Develop Vision and Strategy"


def test_external_id_is_stable_across_variants(dataset):
    """Hierarchy ID 会漂移，PCF ID 不会——这是跨变体关联必须用 external_id 的证据。"""
    ci = {row["external_id"]: row for row in pcf_data.category_rows("APQC_PCF")}
    ce = {row["external_id"]: row for row in pcf_data.category_rows("APQC_PCF_CE")}
    assert set(ci) <= set(ce)
    drifted = [
        external_id for external_id in ci if ci[external_id]["code"] != ce[external_id]["code"]
    ]
    assert drifted, "消费电子版应存在 Hierarchy ID 漂移"


def test_committed_metrics_are_unique_and_linked(dataset):
    metrics = dataset["metrics"]
    assert len(metrics) == 3453
    keys = {(row["external_id"], row["metric_id"]) for row in metrics}
    assert len(keys) == len(metrics)
    assert dataset["metric_conflict_count"] == 138

    known = {
        row["external_id"] for framework in dataset["frameworks"] for row in framework["categories"]
    }
    assert all(row["external_id"] in known for row in metrics)
    assert any(row["name"] for row in metrics)
